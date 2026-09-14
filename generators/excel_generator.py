import pandas as pd
from typing import Dict, List, Any, Optional
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime


# ---------------------------------------------------------------------------
# COSTANTI — Colori per stati e priorità
# ---------------------------------------------------------------------------
STATE_COLORS = {
    "OK": PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    "WARN": PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    "FAIL": PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    "INFO": PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid"),
    "N/A": PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"),
}

PRIORITY_COLORS = {
    1: PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    2: PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    3: PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
}

HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

SECTION_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
SECTION_FONT = Font(bold=True, size=12)

# Larghezze colonne predefinite per tipo di foglio
COLUMN_WIDTHS = {
    "AUDIT": [10, 12, 30, 8, 10, 50, 40, 40],
    "Checklist": [5, 15, 25, 12, 40, 12, 45, 15, 8, 35, 25],
    "drilldown": [60, 60, 60, 60],
    "Focus Keywords": [60, 30, 60],
    "GA4_Landing": [40, 12, 12, 12, 12, 15],
    "GA4_Exit": [40, 12, 12, 12, 12],
    "GA4_Geo": [40, 15],
}


# ---------------------------------------------------------------------------
# HELPER 1 — Creazione DataFrame drill-down
# ---------------------------------------------------------------------------
def create_drilldown_dataframe(
    data: List[Dict],
    columns: List[str],
    empty_message: str = "Nessun problema rilevato"
) -> pd.DataFrame:
    """Crea un DataFrame per fogli drill-down, gestendo il caso vuoto."""
    if not data:
        na_row = {col: "N/A" for col in columns}
        na_row[columns[-1]] = empty_message
        return pd.DataFrame([na_row], columns=columns)

    return pd.DataFrame(data, columns=columns)


# ---------------------------------------------------------------------------
# HELPER 2 — Formattazione header
# ---------------------------------------------------------------------------
def apply_header_format(worksheet):
    for cell in worksheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT


# ---------------------------------------------------------------------------
# HELPER 3 — Larghezza colonne automatica
# ---------------------------------------------------------------------------
def auto_adjust_columns(worksheet, max_width: int = 80):
    for col in worksheet.columns:
        max_length = 0
        column_letter = col[0].column_letter

        for cell in col:
            try:
                cell_len = len(str(cell.value)) if cell.value else 0
                if cell_len > max_length:
                    max_length = cell_len
            except:
                pass

        adjusted_width = min(max_length + 2, max_width)
        worksheet.column_dimensions[column_letter].width = adjusted_width


# ---------------------------------------------------------------------------
# HELPER 4 — Imposta larghezze colonne da lista
# ---------------------------------------------------------------------------
def set_column_widths(worksheet, widths: List[int]):
    for i, width in enumerate(widths, 1):
        worksheet.column_dimensions[get_column_letter(i)].width = width


# ---------------------------------------------------------------------------
# HELPER 5 — Applica colori stato a colonna specifica
# ---------------------------------------------------------------------------
def apply_state_colors(worksheet, column_index: int, color_map: Dict = None):
    if color_map is None:
        color_map = STATE_COLORS

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        cell = row[column_index - 1]
        if cell.value in color_map:
            cell.fill = color_map[cell.value]
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(bold=True)


# ---------------------------------------------------------------------------
# HELPER 6 — Applica colori priorità a colonna specifica
# ---------------------------------------------------------------------------
def apply_priority_colors(worksheet, column_index: int, color_map: Dict = None):
    if color_map is None:
        color_map = PRIORITY_COLORS

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        cell = row[column_index - 1]
        if cell.value in color_map:
            cell.fill = color_map[cell.value]
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(bold=True)


# ---------------------------------------------------------------------------
# HELPER 7 — Applica wrap text a tutte le celle dati
# ---------------------------------------------------------------------------
def apply_wrap_text(worksheet, start_row: int = 2):
    for row in worksheet.iter_rows(min_row=start_row, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


# ---------------------------------------------------------------------------
# HELPER 8 — Estrazione dati drill-down con mapping campi
# ---------------------------------------------------------------------------
def extract_drilldown_data(
    drilldown: Dict,
    key: str,
    columns: List[str],
    field_mapping: Dict[str, str]
) -> List[Dict]:
    data = drilldown.get(key, [])
    if not data:
        return []

    result = []
    for item in data:
        row = {}
        for col in columns:
            field = field_mapping.get(col, col)
            row[col] = item.get(field, "N/A")
        result.append(row)

    return result


class ExcelGenerator:
    """Genera il report Excel dell'audit SEO nel formato target."""

    def __init__(self):
        self.state_colors = STATE_COLORS
        self.priority_colors = PRIORITY_COLORS

    def generate(self, processed_data: Dict, output_path: str, raw_data: Dict = None):
        """Genera il file Excel completo."""

        summary = processed_data['summary']
        audit_rows = processed_data['audit']
        checklist_rows = processed_data['checklist']
        drilldown = processed_data.get('drilldown', {})

        if raw_data is None:
            raw_data = {}

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # 1. Executive Summary
            self._write_summary(writer, summary, audit_rows)

            # 2. AUDIT
            self._write_audit(writer, audit_rows)

            # 3. Checklist principale
            self._write_checklist(writer, checklist_rows, audit_rows)

            # 4. Fogli drill-down
            self._write_drilldown_sheet(
                writer, 'HTML - IMG', drilldown, 'images',
                columns=['URl', 'problem'],
                field_mapping={'URl': 'url', 'problem': 'problem'}
            )

            self._write_drilldown_sheet(
                writer, 'Title', drilldown, 'titles',
                columns=['URl', 'meta_title', 'problem'],
                field_mapping={'URl': 'url', 'meta_title': 'meta_title', 'problem': 'problem'}
            )

            self._write_drilldown_sheet(
                writer, 'Description', drilldown, 'descriptions',
                columns=['URl', 'meta_description', 'problem'],
                field_mapping={'URl': 'url', 'meta_description': 'meta_description', 'problem': 'problem'}
            )

            self._write_drilldown_sheet(
                writer, 'HTML - HEADINGS', drilldown, 'headings',
                columns=['URl', 'h1-1', 'h1-2', 'problem'],
                field_mapping={'URl': 'url', 'h1-1': 'h1-1', 'h1-2': 'h1-2', 'problem': 'problem'}
            )

            # 4b. Focus Keywords (v2.3.8)
            self._write_focus_keywords_sheet(writer, drilldown)

            # 5. Fogli GA4 avanzati (se raw_data è disponibile)
            ga4_data = raw_data.get('ga4', {})
            if ga4_data:
                self._write_ga4_landing_pages(writer, ga4_data)
                self._write_ga4_exit_pages(writer, ga4_data)
                self._write_ga4_geo(writer, ga4_data)

            # 6. Consigli generali
            self._write_consigli(writer, audit_rows)

        # Formatta il file
        self._format_excel(output_path)

    # ==================================================================
    # SCRITTURA FOGLI
    # ==================================================================

    def _write_drilldown_sheet(
        self,
        writer,
        sheet_name: str,
        drilldown: Dict,
        data_key: str,
        columns: List[str],
        field_mapping: Dict[str, str]
    ):
        """Scrive un foglio drill-down generico."""
        data = extract_drilldown_data(drilldown, data_key, columns, field_mapping)
        df = create_drilldown_dataframe(data, columns)
        df.to_excel(writer, sheet_name=sheet_name, index=False)

    def _write_focus_keywords_sheet(self, writer, drilldown: Dict):
        """Scrive il foglio Focus Keywords (v2.3.8).

        Elenca per ogni URL con keyword assegnata in focus_keywords.yaml i
        problemi rilevati: keyword non in title, non in H1, density fuori
        range. La lista è già deduplicata dal processor.
        """
        focus_data = drilldown.get('focus_keywords', [])

        if not focus_data:
            rows = [{
                'URL': 'N/A',
                'Focus Keyword': 'N/A',
                'Problem': 'Nessun problema rilevato o focus_keywords.yaml non configurato',
            }]
        else:
            rows = []
            for item in focus_data:
                rows.append({
                    'URL': item.get('url', ''),
                    'Focus Keyword': item.get('focus_keyword', ''),
                    'Problem': item.get('problem', ''),
                })

        df = pd.DataFrame(rows)
        df.to_excel(writer, sheet_name='Focus Keywords', index=False)

    def _write_ga4_landing_pages(self, writer, ga4_data: Dict):
        """Scrive il foglio GA4 Landing Pages."""
        landing_pages = ga4_data.get('landing_pages', {}).get('landing_pages', [])

        if not landing_pages:
            rows = [{'Pagina': 'N/A', 'Sessioni': 'N/A', 'Utenti': 'N/A',
                     'Engaged': 'N/A', 'Bounce': 'N/A', 'Durata': 'N/A'}]
        else:
            rows = []
            for lp in landing_pages[:15]:
                rows.append({
                    'Pagina': lp.get('page', ''),
                    'Sessioni': lp.get('sessions', 0),
                    'Utenti': lp.get('users', 0),
                    'Engaged': lp.get('engaged_sessions', 0),
                    'Bounce': f"{lp.get('bounce_rate', 0)*100:.1f}%",
                    'Durata': f"{lp.get('avg_duration', 0)/60:.1f}m",
                })

        df = pd.DataFrame(rows)
        df.to_excel(writer, sheet_name='GA4 Landing Pages', index=False)

    def _write_ga4_exit_pages(self, writer, ga4_data: Dict):
        """Scrive il foglio GA4 Exit Pages."""
        exit_pages = ga4_data.get('exit_pages', {}).get('exit_pages', [])

        if not exit_pages:
            rows = [{'Pagina': 'N/A', 'Exits (stima)': 'N/A', 'Views': 'N/A',
                     'Sessioni': 'N/A', 'Bounce': 'N/A'}]
        else:
            rows = []
            for ep in exit_pages[:15]:
                rows.append({
                    'Pagina': ep.get('page', ''),
                    'Exits (stima)': ep.get('estimated_exits', 0),
                    'Views': ep.get('views', 0),
                    'Sessioni': ep.get('sessions', 0),
                    'Bounce': f"{ep.get('bounce_rate', 0)*100:.1f}%"
                })

        df = pd.DataFrame(rows)
        df.to_excel(writer, sheet_name='GA4 Exit Pages', index=False)

    def _write_ga4_geo(self, writer, ga4_data: Dict):
        """Scrive il foglio GA4 Geo Distribution."""
        geo_data = ga4_data.get('geo_distribution', {}).get('geo_distribution', [])

        if not geo_data:
            rows = [{'Paese': 'N/A', 'Sessioni': 'N/A'}]
        else:
            rows = []
            for geo in geo_data[:10]:
                rows.append({
                    'Paese': geo.get('country', ''),
                    'Sessioni': geo.get('sessions', 0)
                })

        df = pd.DataFrame(rows)
        df.to_excel(writer, sheet_name='GA4 Geo', index=False)

    def _write_summary(self, writer, summary: Dict, audit_rows: List[Dict]):
        """Scrive il foglio Executive Summary."""
        by_category = self._group_by_category(audit_rows)
        rows = self._build_summary_rows(summary, by_category)

        df = pd.DataFrame(rows, columns=["Metrica", "Dettaglio"])
        df.to_excel(writer, sheet_name="Executive Summary", index=False)

    def _group_by_category(self, audit_rows: List[Dict]) -> Dict[str, Dict[str, int]]:
        """Raggruppa i check per categoria e stato."""
        by_category = {}
        for row in audit_rows:
            cat = row['Categoria']
            if cat not in by_category:
                by_category[cat] = {'OK': 0, 'WARN': 0, 'FAIL': 0, 'INFO': 0, 'N/A': 0}
            by_category[cat][row['Stato']] += 1
        return by_category

    def _build_summary_rows(self, summary: Dict, by_category: Dict) -> List[List]:
        """Costruisce le righe del summary."""
        rows = [
            ["DATI CLIENTE", ""],
            ["Dominio", summary['domain']],
            ["Data Audit", summary['date']],
            ["Versione Report", summary['version']],
            ["", ""],
            ["HEALTH SCORE", ""],
            ["Punteggio Tecnico", f"{summary['health_score']}/100"],
            ["", ""],
            ["STATISTICHE", ""],
            ["Totale check eseguiti", summary['total_checks']],
            ["Errori critici (FAIL)", summary['fails']],
            ["Warning (WARN)", summary['warnings']],
            ["OK", summary['oks']],
            ["Info (INFO)", summary.get('infos', 0)],
            ["Non applicabili (N/A)", summary.get('na', 0)],
            ["", ""],
            ["TOP 3 CRITICITÀ", ""],
        ]

        top_criticità = summary.get('top_criticita', [])[:3]
        if top_criticità:
            for i, crit in enumerate(top_criticità, 1):
                rows.append([f"{i}.", f"{crit['Elemento Analizzato']}: {crit['Risultato / Evidenza']}"])
        else:
            rows.append(["", "Nessuna criticità rilevata!"])

        rows.extend([
            ["", ""],
            ["TOP PUNTI DI FORZA", ""],
        ])

        for i, pf in enumerate(summary.get('top_punti_forza', [])[:3], 1):
            rows.append([f"{i}.", pf['Elemento Analizzato']])

        rows.extend([
            ["", ""],
            ["RIPARTIZIONE PER CATEGORIA", ""],
            ["Categoria", "OK / WARN / FAIL / INFO / N/A"],
        ])

        for cat, counts in sorted(by_category.items()):
            stats = f"{counts['OK']} / {counts['WARN']} / {counts['FAIL']} / {counts['INFO']} / {counts['N/A']}"
            rows.append([cat, stats])

        return rows

    def _write_audit(self, writer, audit_rows: List[Dict]):
        """Scrive il foglio AUDIT."""
        canonical = [
            "ID Audit", "Categoria", "Elemento Analizzato", "Stato", "Severità",
            "Risultato / Evidenza", "URL / Link Evidenza", "Note Tecniche"
        ]

        if not audit_rows:
            df = pd.DataFrame(columns=canonical)
        else:
            all_keys = set()
            for row in audit_rows:
                all_keys.update(row.keys())

            columns = [c for c in canonical if c in all_keys]
            extra = [k for k in all_keys if k not in columns]
            columns += sorted(extra)

            df = pd.DataFrame(audit_rows, columns=columns)

        df.to_excel(writer, sheet_name="AUDIT", index=False)

    def _write_checklist(self, writer, checklist_rows: List[Dict], audit_rows: List[Dict]):
        """Scrive il foglio Checklist."""
        audit_map = {row['ID Audit']: row for row in audit_rows}

        rows = []
        for check_num, chk in enumerate(checklist_rows, 1):
            audit_id = chk['Rif. Audit']
            audit_row = audit_map.get(audit_id, {})

            rows.append({
                '#': check_num,
                'Category': audit_row.get('Categoria', ''),
                'Activities': audit_row.get('Elemento Analizzato', ''),
                'Fase': chk.get('Fase', ''),
                'Results': audit_row.get('Risultato / Evidenza', ''),
                'Priority (1 alta, 3 bassa)': chk.get('Priorità', 2),
                'to do': chk.get('Azione Richiesta', ''),
                'Owner': chk.get('Owner', ''),
                'Check': 'FALSE',
                'Note': audit_row.get('Note Tecniche', ''),
                'NOTE PF': ''
            })

        df = pd.DataFrame(rows)
        df.to_excel(writer, sheet_name='Checklist', index=False)

    def _write_consigli(self, writer, audit_rows: List[Dict]):
        """Scrive il foglio Consigli generali."""
        rows = self._generate_consigli_rows(audit_rows)
        df = pd.DataFrame(rows)
        df.to_excel(writer, sheet_name='Consigli generali', index=False)

    def _generate_consigli_rows(self, audit_rows: List[Dict]) -> List[Dict]:
        """Genera le righe del foglio Consigli generali dai FAIL/WARN."""
        issues_by_cat: Dict[str, List[Dict]] = {}
        for row in audit_rows:
            if row.get('Stato') in ('FAIL', 'WARN'):
                cat = row.get('Categoria', 'Altro')
                issues_by_cat.setdefault(cat, []).append(row)

        if not issues_by_cat:
            return [{'Consigli generali da applicare':
                     '✓ Nessuna criticità rilevata: il sito risulta ben ottimizzato.'}]

        def cat_priority(cat: str):
            issues = issues_by_cat[cat]
            has_fail = any(i.get('Stato') == 'FAIL' for i in issues)
            min_sev = min((i.get('Severità', 3) for i in issues), default=3)
            return (0 if has_fail else 1, min_sev, cat)

        sorted_cats = sorted(issues_by_cat.keys(), key=cat_priority)

        rows: List[Dict] = []
        rows.append({'Consigli generali da applicare':
                     'PRIORITÀ DI INTERVENTO (ordine consigliato)'})
        rows.append({'Consigli generali da applicare': ''})

        for cat in sorted_cats:
            issues = issues_by_cat[cat]
            fails = sorted(
                [i for i in issues if i.get('Stato') == 'FAIL'],
                key=lambda x: x.get('Severità', 3)
            )
            warns = sorted(
                [i for i in issues if i.get('Stato') == 'WARN'],
                key=lambda x: x.get('Severità', 3)
            )

            header = f"■ {cat}"
            parts = []
            if fails:
                parts.append(f"{len(fails)} FAIL")
            if warns:
                parts.append(f"{len(warns)} WARN")
            if parts:
                header += " — " + ", ".join(parts)
            rows.append({'Consigli generali da applicare': header})

            for issue in fails + warns:
                symbol = '✗' if issue.get('Stato') == 'FAIL' else '⚠'
                element = issue.get('Elemento Analizzato', '')
                note = (issue.get('Note Tecniche') or '').strip()
                if note:
                    rows.append({
                        'Consigli generali da applicare': f"  {symbol} {element}: {note}"
                    })
                else:
                    rows.append({
                        'Consigli generali da applicare': f"  {symbol} {element}"
                    })

            rows.append({'Consigli generali da applicare': ''})

        return rows

    # ==================================================================
    # FORMATTAZIONE
    # ==================================================================

    def _format_excel(self, file_path: str):
        """Applica la formattazione al file Excel."""
        wb = load_workbook(file_path)

        formatters = {
            "Executive Summary": self._format_summary_sheet,
            "AUDIT": self._format_audit_sheet,
            "Checklist": self._format_checklist_sheet,
            "HTML - IMG": self._format_drilldown_sheet,
            "Title": self._format_drilldown_sheet,
            "Description": self._format_drilldown_sheet,
            "HTML - HEADINGS": self._format_drilldown_sheet,
            "Focus Keywords": self._format_focus_keywords_sheet,
            "GA4 Landing Pages": self._format_ga4_landing_sheet,
            "GA4 Exit Pages": self._format_ga4_exit_sheet,
            "GA4 Geo": self._format_ga4_geo_sheet,
            "Consigli generali": self._format_consigli_sheet,
        }

        for sheet_name, formatter in formatters.items():
            if sheet_name in wb.sheetnames:
                formatter(wb[sheet_name])

        wb.save(file_path)

    def _format_summary_sheet(self, worksheet):
        """Formatta il foglio Executive Summary."""
        worksheet.column_dimensions['A'].width = 30
        worksheet.column_dimensions['B'].width = 60

        for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row):
            if row[0].value == "Punteggio Tecnico" and len(row) >= 2:
                row[1].font = Font(bold=True, size=16, color="0000FF")
                break

        section_titles = {
            "DATI CLIENTE", "HEALTH SCORE", "STATISTICHE",
            "TOP 3 CRITICITÀ", "TOP PUNTI DI FORZA",
            "RIPARTIZIONE PER CATEGORIA"
        }

        for row in worksheet.iter_rows(min_row=1, max_row=worksheet.max_row):
            cell_a = row[0]
            if cell_a.value in section_titles:
                cell_a.font = SECTION_FONT
                cell_a.fill = SECTION_FILL

    def _format_audit_sheet(self, worksheet):
        """Formatta il foglio AUDIT."""
        set_column_widths(worksheet, COLUMN_WIDTHS["AUDIT"])
        apply_header_format(worksheet)
        apply_state_colors(worksheet, column_index=4)
        apply_wrap_text(worksheet)

    def _format_checklist_sheet(self, worksheet):
        """Formatta il foglio Checklist."""
        set_column_widths(worksheet, COLUMN_WIDTHS["Checklist"])
        apply_header_format(worksheet)
        apply_priority_colors(worksheet, column_index=6)
        apply_wrap_text(worksheet)

    def _format_drilldown_sheet(self, worksheet):
        """Formatta un foglio drill-down generico."""
        apply_header_format(worksheet)
        auto_adjust_columns(worksheet)
        apply_wrap_text(worksheet)

    def _format_focus_keywords_sheet(self, worksheet):
        """Formatta il foglio Focus Keywords (v2.3.8)."""
        set_column_widths(worksheet, COLUMN_WIDTHS["Focus Keywords"])
        apply_header_format(worksheet)
        apply_wrap_text(worksheet)

    def _format_ga4_landing_sheet(self, worksheet):
        """Formatta il foglio GA4 Landing Pages."""
        set_column_widths(worksheet, COLUMN_WIDTHS["GA4_Landing"])
        apply_header_format(worksheet)
        apply_wrap_text(worksheet)

    def _format_ga4_exit_sheet(self, worksheet):
        """Formatta il foglio GA4 Exit Pages."""
        set_column_widths(worksheet, COLUMN_WIDTHS["GA4_Exit"])
        apply_header_format(worksheet)
        apply_wrap_text(worksheet)

    def _format_ga4_geo_sheet(self, worksheet):
        """Formatta il foglio GA4 Geo."""
        set_column_widths(worksheet, COLUMN_WIDTHS["GA4_Geo"])
        apply_header_format(worksheet)
        apply_wrap_text(worksheet)

    def _format_consigli_sheet(self, worksheet):
        """Formatta il foglio Consigli generali."""
        worksheet.column_dimensions['A'].width = 120
        apply_header_format(worksheet)
        apply_wrap_text(worksheet)